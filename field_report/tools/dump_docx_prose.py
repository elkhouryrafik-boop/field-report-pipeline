#!/usr/bin/env python
"""Dump every paragraph and table cell of a docx to plain text.

Feeds the field-report scan gate:
    python -m field_report.tools.dump_docx_prose <docx> <out.txt>
    python -m field_report.tools.dump_docx_prose <docx> <out.txt>
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document


def main(argv: list[str]) -> int:
    if len(argv) != 2:
        print(__doc__, file=sys.stderr)
        return 2
    src, dst = Path(argv[0]), Path(argv[1])
    doc = Document(str(src))
    lines: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(p.text.strip())
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                text = c.text.strip()
                if text:
                    lines.append(text)
    dst.write_text("\n\n".join(lines), encoding="utf-8")
    print(f"{len(lines)} blocks -> {dst}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
