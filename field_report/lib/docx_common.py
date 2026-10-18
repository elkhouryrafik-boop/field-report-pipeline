"""Shared python-docx / openpyxl helpers for the report (s7), the reviewer sheet (s9) and the quantities sheet (s8).

Everything profile-dependent (font, scale) is passed in; nothing here knows a client.
"""
from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

CREATOR = "field-report-pipeline"
MARKER_RE = re.compile(r"(\[to be (?:confirmed|verified)[^\]]*\])")
HEADING_COLOR = "1F3864"
HEADER_FILL = "1F3864"
BAND_FILL = "F2F2F2"
MARKER_COLOR = "9E1B1B"
# One fill per rank on the profile scale, most severe first; N/A is grey.
RANK_FILLS = ["F4CCCC", "FCE5CD", "FFF2CC", "D9EAD3", "DDEBF7"]
NA_FILL = "EDEDED"


# --------------------------------------------------------------------------- metadata

def neutral_core_properties(doc, title: str, subject: str = "") -> None:
    """creator / lastModifiedBy = field-report-pipeline; no company, no template path."""
    cp = doc.core_properties
    now = datetime.now(timezone.utc).replace(microsecond=0, tzinfo=None)
    cp.author = CREATOR
    cp.last_modified_by = CREATOR
    cp.title = title
    cp.subject = subject
    cp.comments = ""
    cp.keywords = ""
    cp.category = ""
    cp.content_status = ""
    cp.created = now
    cp.modified = now
    cp.revision = 1


def neutral_workbook_properties(wb, title: str) -> None:
    wb.properties.creator = CREATOR
    wb.properties.lastModifiedBy = CREATOR
    wb.properties.title = title
    wb.properties.description = None
    wb.properties.keywords = None
    wb.properties.category = None
    wb.properties.subject = None


APP_XML = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
           '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
           '<Application>field-report-pipeline</Application></Properties>')


def save_docx(doc, path: Path) -> None:
    """Save, then replace docProps/app.xml: python-docx's default template leaves a template
    name ("Normal.dotm") and Word statistics there. Nothing about the machine or a template ships."""
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    tmp = Path(path).with_suffix(".tmp")
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/app.xml":
                data = APP_XML.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


# --------------------------------------------------------------------------- document

def new_document(prof: dict, title: str, *, landscape: bool = False, margin_in: float = 0.75):
    """A blank document with the profile's body font and neutral metadata."""
    doc = Document()
    font = prof["font"]
    normal = doc.styles["Normal"]
    normal.font.name = font["name"]
    normal.font.size = Pt(font["size_pt"])
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for slot in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(slot), font["name"])
    for level, size in ((1, 16), (2, 13), (3, 12)):
        st = doc.styles[f"Heading {level}"]
        st.font.name = font["name"]
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(HEADING_COLOR)
        hr = st.element.get_or_add_rPr()
        hf = hr.find(qn("w:rFonts"))
        if hf is None:
            hf = OxmlElement("w:rFonts")
            hr.append(hf)
        for slot in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            hf.set(qn(slot), font["name"])
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if hf.get(qn(attr)) is not None:
                del hf.attrib[qn(attr)]
    sec = doc.sections[0]
    if landscape:
        sec.page_width, sec.page_height = Inches(11), Inches(8.5)
    else:
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(margin_in))
    neutral_core_properties(doc, title)
    return doc


def content_width_in(doc) -> float:
    sec = doc.sections[0]
    return (sec.page_width - sec.left_margin - sec.right_margin) / 914400


def shade_cell(cell, hex_fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn("w:shd")):
        tcpr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def set_col_widths(table, widths_in: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for cell, w in zip(row.cells, widths_in):
            cell.width = Inches(w)


def write_marked_runs(p, text: str, *, bold: bool = False, italic: bool = False,
                      color: str | None = None, size: float | None = None) -> None:
    """Add `text` to paragraph `p`; any [to be confirmed ...] marker is bold red."""
    for part in MARKER_RE.split(text):
        if not part:
            continue
        run = p.add_run(part)
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        if MARKER_RE.fullmatch(part):
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(MARKER_COLOR)
        elif color:
            run.font.color.rgb = RGBColor.from_string(color)


def set_cell(cell, text: str, *, bold: bool = False, italic: bool = False, size: float | None = None,
             color: str | None = None, fill: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    write_marked_runs(p, text, bold=bold, italic=italic, color=color, size=size)
    if fill:
        shade_cell(cell, fill)


def add_table(doc, header: list[str] | None, n_rows: int, n_cols: int, widths_in: list[float] | None = None,
              size: float = 10):
    """A grid table; the header row (if any) is shaded and repeats on each page."""
    table = doc.add_table(rows=n_rows + (1 if header else 0), cols=n_cols)
    table.style = "Table Grid"
    if header:
        for c, text in zip(table.rows[0].cells, header):
            set_cell(c, text, bold=True, color="FFFFFF", fill=HEADER_FILL, size=size)
        trpr = table.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trpr.append(th)
    if widths_in:
        set_col_widths(table, widths_in)
    return table


def rating_fill(value: str, scale: list[str]) -> str:
    if value in scale:
        return RANK_FILLS[min(scale.index(value), len(RANK_FILLS) - 1)]
    return NA_FILL


def add_body(doc, text: str, *, italic: bool = False, bold: bool = False, size: float | None = None,
             color: str | None = None, align=None, space_after: float = 6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    write_marked_runs(p, text, bold=bold, italic=italic, size=size, color=color)
    return p


def page_break(doc) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# --------------------------------------------------------------------------- TOC

TOC_INSTR = ' TOC \\o "1-2" \\h \\z \\u '


def _fld(kind: str):
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), kind)
    r = OxmlElement("w:r")
    r.append(el)
    return r


def insert_toc(doc, anchor, entries: list[tuple[int, str]]) -> None:
    """Splice a Word TOC field in right after paragraph `anchor`.

    The field result is a static list of the headings this build emitted, so a reader
    whose Word never refreshes fields still gets a usable contents page (no page numbers).
    """
    made = []
    title = doc.add_paragraph()
    write_marked_runs(title, "Contents", bold=True, size=16, color=HEADING_COLOR)
    made.append(title)
    body = entries or [(1, "(no sections)")]
    for i, (level, text) in enumerate(body):
        p = doc.add_paragraph()
        if level >= 2:
            p.paragraph_format.left_indent = Inches(0.28)
        if i == 0:
            p._p.append(_fld("begin"))
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = TOC_INSTR
            r = OxmlElement("w:r")
            r.append(instr)
            p._p.append(r)
            p._p.append(_fld("separate"))
        write_marked_runs(p, text)
        if i == len(body) - 1:
            p._p.append(_fld("end"))
        made.append(p)
    brk = doc.add_paragraph()
    brk.add_run().add_break(WD_BREAK.PAGE)
    made.append(brk)
    for p in reversed(made):
        anchor._p.addnext(p._p)


# --------------------------------------------------------------------------- images

def photo_stream(path: Path, max_px: int = 1000) -> io.BytesIO:
    """A re-encoded JPEG copy: bounded size, EXIF never carried over."""
    from PIL import Image

    with Image.open(path) as im:
        im = im.convert("RGB")
        im.thumbnail((max_px, max_px))
        buf = io.BytesIO()
        im.save(buf, format="JPEG", quality=82)   # no exif= argument: metadata is dropped
    buf.seek(0)
    return buf


def docx_text(doc) -> str:
    """Every paragraph and table cell, in reading order, as plain text."""
    parts: list[str] = []
    body = doc.element.body
    for node in body.iter():
        if node.tag == qn("w:p"):
            parts.append("".join(t.text or "" for t in node.iter(qn("w:t"))))
    return "\n".join(parts)
