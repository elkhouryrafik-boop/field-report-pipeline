#!/usr/bin/env python
"""s8_build_quantities.py - stage 8: findings.json -> quantities.xlsx (profiles with `quantities_sheet`).

    python -m field_report.stages.s8_build_quantities --config <config> [--run-dir DIR] [--force]

One row per spoken quantity: category, asset, unit, number, the words that were said, whether a
reviewer has confirmed it, and where it came from. Quantities only - there are no price, rate or
total columns, and the gate fails the file if one appears. Unconfirmed figures are highlighted.

A profile without `quantities_sheet: true` (site-inspection, event-log) writes nothing and exits 0.
Exit codes: 0 ok / not applicable, 2 bad invocation, 4 failed, 5 missing upstream.
"""
from __future__ import annotations

import argparse
import re
import sys
import traceback
from pathlib import Path

from field_report.lib import common, profiles
from field_report.lib import docx_common as dc


STAGE = "s8_build_quantities"
OUT_NAME = "quantities.xlsx"
HEADER = ["Item", "Category", "Asset", "Unit", "Quantity", "Said as", "Status", "Stop", "Source"]
COL_WIDTHS = [10, 22, 30, 8, 12, 30, 16, 10, 34]
PRICE_WORDS = re.compile(r"\b(rate|price|unit price|amount|total|cost)\b", re.I)
HEADER_FILL, OPEN_FILL = "1F3864", "FFF2CC"


def parse_args(argv=None):
    p = argparse.ArgumentParser(description="stage 8: build quantities.xlsx")
    p.add_argument("--config", required=True)
    p.add_argument("--run-dir", default=None)
    p.add_argument("--force", action="store_true")
    p.add_argument("--log-level", default="INFO", choices=["DEBUG", "INFO", "WARNING", "ERROR"])
    return p.parse_args(argv)


def build_rows(findings: dict) -> list[dict]:
    rows = []
    for sector in findings["sectors"]:
        if not sector.get("present", True):
            continue
        for a in sector["assets"]:
            for i, q in enumerate(a.get("quantities", [])):
                refs = [e["ref"] for e in a.get("source_evidence", []) if e.get("type") == "transcript"]
                rows.append({
                    "item": f"{a['id']}/{i + 1}", "category": sector["sector_name"], "asset": a["name"],
                    "unit": "" if q["unit"] == "other" else q["unit"],
                    "number": q.get("value"), "text": q.get("value_text"),
                    "said": q.get("verbatim") or "", "confirmed": bool(q.get("confirmed")),
                    "stop": a.get("map_marker") or "", "source": ", ".join(refs) or "field team"})
    return rows


def write_workbook(rows: list[dict], findings: dict, cfg: dict, out_path: Path, synthetic: bool) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Quantities"
    ws.append(HEADER)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEADER_FILL)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        status = "Confirmed" if r["confirmed"] else "To be confirmed"
        number = r["number"] if r["number"] is not None else (r["text"] or "")
        ws.append([r["item"], r["category"], r["asset"], r["unit"], number, r["said"], status,
                   r["stop"], r["source"]])
        if not r["confirmed"]:
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor=OPEN_FILL)
    for col, w in zip("ABCDEFGHI", COL_WIDTHS):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    notes = wb.create_sheet("Notes")
    lines = [f"Site: {cfg['site']}",
             f"Profile: {findings['profile']}",
             "Quantities only, as spoken in the field voice notes. Highlighted rows are not yet confirmed.",
             "Nothing here is measured or priced."]
    if synthetic:
        lines.append("Generated from synthetic sample data.")
    for ln in lines:
        notes.append([ln])
    notes.column_dimensions["A"].width = 100
    dc.neutral_workbook_properties(wb, f"Quantities - {cfg['site']}")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))


def _run(cfg, run_dir: Path, log, args) -> int:
    prof = profiles.profile(cfg)
    if not prof.get("quantities_sheet"):
        print(f"SKIP {STAGE}: profile {prof['name']!r} has no quantities sheet")
        return 0
    findings = common.read_stage(run_dir, "findings", "findings")
    out = run_dir / OUT_NAME
    fpath = run_dir / "findings.json"
    if not args.force and out.is_file() and out.stat().st_mtime >= fpath.stat().st_mtime:
        print(f"SKIP {STAGE}: up to date")
        return 0
    rows = build_rows(findings)
    from field_report.stages.s7_build_report import is_synthetic
    write_workbook(rows, findings, cfg, out, is_synthetic(cfg))
    log.info("quantities: %d rows -> %s", len(rows), out)
    return 0


def main(argv=None) -> int:
    args = parse_args(argv)
    try:
        cfg = common.load_config(args.config)
    except common.ConfigError as e:
        sys.stderr.write(f"CONFIG ERROR: {e}\n")
        return 2
    run_dir = common.resolve_run_dir(cfg, args.run_dir)
    log = common.get_logger(run_dir, STAGE, args.log_level)
    try:
        return _run(cfg, run_dir, log, args)
    except common.MissingUpstream as e:
        log.error("missing upstream: %s", e)
        return 5
    except ImportError as e:
        sys.stderr.write(f"missing dependency: {e} (pip install openpyxl)\n")
        return 2
    except Exception:  # noqa: BLE001
        tb = traceback.format_exc()
        (run_dir / "logs" / f"{STAGE}.err").write_text(tb, encoding="utf-8")
        log.error("unhandled exception:\n%s", tb)
        return 4


if __name__ == "__main__":
    sys.exit(main())
