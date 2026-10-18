#!/usr/bin/env python
"""s7_build_report.py - stage 7: findings.json -> report.docx (profile-driven).

    python -m field_report.stages.s7_build_report --config <config> [--run-dir DIR] [--force]

The structure comes from the profile (`report` wording, categories, rating scale, font):

  cover -> contents -> summary -> stops and map -> findings (by category; or a chronological
  log for a profile that does not classify) -> priorities -> open items -> voice-note log

Rules the gate (s10) re-checks:
  * every photo in findings.json is embedded exactly once, with its caption;
  * every quantity is printed with the words that were actually said, and an unconfirmed one
    carries a bold-red "[to be confirmed: quantity]" marker that the reviewer sheet mirrors;
  * nothing the field team has not looked at is presented as verified - an auto-accepted
    run says "unreviewed" on the cover, in the summary and on every finding;
  * no individual is named (the people who did the work are "the field team").

When `sample_manifest.json` sits next to the config, the cover says "Generated from synthetic
sample data". The map is a PNG (`map_image` in the config); there is no PDF map input.

Writes report.docx and report_build.json (what was embedded, how many markers).
Exit codes: 0 ok, 2 bad invocation, 3 partial (a photo file could not be embedded), 4 failed,
5 missing upstream.
"""
from __future__ import annotations

import argparse
import json
import sys
import traceback
from datetime import datetime
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from field_report.lib import common, profiles
from field_report.lib import docx_common as dc
from field_report.lib import open_items
from field_report.lib.privacy import FIELD_TEAM

STAGE = "s7_build_report"
OUT_NAME = "report.docx"
BUILD_NAME = "report_build.json"
PHOTOS_PER_ROW = 3
SYNTHETIC_LINE = "Generated from synthetic sample data"


def parse_args(argv=None):
    p = argparse.ArgumentParser(description="stage 7: build report.docx")
    p.add_argument("--config", required=True)
    p.add_argument("--run-dir", default=None)
    p.add_argument("--force", action="store_true")
    p.add_argument("--log-level", default="INFO", choices=["DEBUG", "INFO", "WARNING", "ERROR"])
    return p.parse_args(argv)


# --------------------------------------------------------------------------- small helpers

def long_date(iso: str) -> str:
    d = datetime.strptime(iso[:10], "%Y-%m-%d")
    return f"{d.day} {d.strftime('%B %Y')}"


def hhmm(ts: str | None) -> str:
    return (ts or "")[11:16]


def is_synthetic(cfg: dict) -> bool:
    """True when the config sits next to a sample_manifest.json written by examples/make_sample.py."""
    return (Path(cfg["_dir"]) / "sample_manifest.json").is_file()


def review_state(findings: dict) -> tuple[int, int]:
    stops = findings.get("stops", [])
    return sum(1 for s in stops if s.get("human_verified")), len(stops)


def review_sentence(findings: dict) -> str:
    done, total = review_state(findings)
    if total and done == total:
        return f"All {total} stops were checked by a reviewer."
    if done == 0:
        return (f"Unreviewed: none of the {total} stops was checked by a person. The transcripts, photo "
                f"captions and ratings below were accepted automatically.")
    return f"{done} of {total} stops were checked by a reviewer; the rest are unreviewed."


def photo_captions(findings: dict) -> dict[str, str]:
    """photo filename -> caption, taken from the evidence the classifier cited."""
    caps: dict[str, str] = {}
    for sector in findings["sectors"]:
        for a in sector.get("assets", []):
            for e in a.get("source_evidence", []):
                if e.get("type") == "photo" and e.get("quote"):
                    caps.setdefault(e["ref"], e["quote"])
    return caps


def transcript_texts(run_dir: Path) -> dict[str, dict]:
    name = "transcripts_en_verified" if (run_dir / "transcripts_en_verified.json").is_file() else "transcripts_en"
    items = common.read_stage(run_dir, name, "transcripts_en")["items"]
    return {i["file"]: i for i in items if i.get("status") in ("ok", "partial")}


# --------------------------------------------------------------------------- photos

class Photos:
    """Embeds each photo at most once, in the order asked, and remembers which."""

    def __init__(self, doc, media_dir: Path, run_dir: Path, captions: dict[str, str]):
        self.doc, self.media_dir, self.run_dir, self.captions = doc, media_dir, run_dir, captions
        self.used: list[str] = []
        self.failed: list[str] = []

    def _path(self, name: str) -> Path:
        redacted = self.run_dir / "media_redacted" / name      # optional face-blurred copy
        return redacted if redacted.is_file() else self.media_dir / name

    def grid(self, names: list[str], width_in: float) -> None:
        names = [n for n in names if n not in self.used]
        for start in range(0, len(names), PHOTOS_PER_ROW):
            row = names[start:start + PHOTOS_PER_ROW]
            table = self.doc.add_table(rows=2, cols=len(row))
            cell_w = width_in / PHOTOS_PER_ROW
            for col, name in enumerate(row):
                number = len(self.used) + 1
                pic_cell, cap_cell = table.rows[0].cells[col], table.rows[1].cells[col]
                pic_cell.width = cap_cell.width = Inches(cell_w)
                try:
                    stream = dc.photo_stream(self._path(name))
                    pic_cell.paragraphs[0].add_run().add_picture(stream, width=Inches(cell_w - 0.12))
                    self.used.append(name)
                except Exception as e:  # noqa: BLE001 - one bad image must not sink the report
                    self.failed.append(f"{name}: {e}")
                    pic_cell.paragraphs[0].add_run(f"[photo {name} could not be embedded]")
                    number = len(self.used) + 1
                caption = self.captions.get(Path(name).stem) or name
                p = cap_cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run(f"Photo {number}: {caption}")
                run.font.size = Pt(8.5)
                run.italic = True
            self.doc.add_paragraph().paragraph_format.space_after = Pt(2)


# --------------------------------------------------------------------------- sections

def add_cover(doc, cfg, prof, findings, synthetic: bool) -> None:
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dc.write_marked_runs(p, prof["report"]["title"], bold=True, size=24, color=dc.HEADING_COLOR)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dc.write_marked_runs(p, profiles.locality_line(cfg), bold=True, size=16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    when = long_date(cfg["date_from"])
    if cfg["date_to"][:10] != cfg["date_from"][:10]:
        when += " to " + long_date(cfg["date_to"])
    dc.write_marked_runs(p, when, size=13)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dc.write_marked_runs(p, f"Prepared by {FIELD_TEAM}", size=12)
    if synthetic:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        dc.write_marked_runs(p, SYNTHETIC_LINE, bold=True, size=12, color=dc.MARKER_COLOR)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dc.write_marked_runs(p, review_sentence(findings), italic=True, size=11)
    dc.page_break(doc)


def add_summary(doc, prof, findings, markers: int, headings: list) -> None:
    rep = prof["report"]
    headings.append((1, "1. " + rep["summary_heading"]))
    doc.add_heading("1. " + rep["summary_heading"], level=1)
    dc.add_body(doc, findings["overall"]["summary"])
    dc.add_body(doc, review_sentence(findings))
    if prof["classify"]:
        scale = profiles.scale_values(prof)
        present = [s for s in findings["sectors"] if s.get("present", True)]
        table = dc.add_table(doc, ["Category", prof["scale"]["label"], "Items"], len(findings["sectors"]), 3,
                             [3.6, 1.6, 1.0])
        for row, s in zip(table.rows[1:], findings["sectors"]):
            dc.set_cell(row.cells[0], s["sector_name"], size=10)
            dc.set_cell(row.cells[1], s["sector_severity"], bold=True, size=10,
                        fill=dc.rating_fill(s["sector_severity"], scale))
            dc.set_cell(row.cells[2], str(len(s["assets"]) if s.get("present", True) else 0), size=10)
        doc.add_paragraph()
        dc.add_body(doc, f"Overall {prof['scale']['label'].lower()}: {findings['overall']['damage_level']}. "
                         f"{len(present)} of {len(findings['sectors'])} categories have at least one item.")
    if markers:
        dc.add_body(doc, f"{markers} figure(s) or photo(s) are still open and flagged "
                         f"in bold red in the text; the reviewer sheet lists each one.")


def add_stops(doc, cfg, prof, findings, headings: list, content_w: float) -> None:
    rep = prof["report"]
    title = f"2. {rep['unit_plural'].capitalize()} and map"
    headings.append((1, title))
    doc.add_heading(title, level=1)
    if cfg.get("map_image"):
        try:
            doc.add_picture(dc.photo_stream(Path(cfg["map_image"]), max_px=1400), width=Inches(min(5.5, content_w)))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            dc.add_body(doc, "Schematic map of the stops (not to scale).", italic=True, size=9,
                        align=WD_ALIGN_PARAGRAPH.CENTER)
        except Exception as e:  # noqa: BLE001 - the map is optional, say so rather than fail
            dc.add_body(doc, f"[the map image could not be embedded: {e}]", italic=True)
    stops = findings["stops"]
    table = dc.add_table(doc, ["#", rep["unit"], "Time", "Position", "Photos", "Reviewed"], len(stops), 6,
                         [0.4, 2.7, 0.7, 1.4, 0.7, 1.1])
    gps = {a["cluster_id"]: a.get("gps") for s in findings["sectors"] for a in s.get("assets", [])
           if a.get("gps")}
    for row, st in zip(table.rows[1:], stops):
        g = gps.get(st["cluster_id"])
        pos = f"{g['lat']:.4f}, {g['lng']:.4f}" + (" (inherited)" if g.get("inherited") else "") if g else "no pin"
        vals = [str(st["seq"]), st["title"], hhmm(st.get("start_ts")), pos, str(len(st["photos"])),
                "yes" if st["human_verified"] else "unreviewed"]
        for c, v in zip(row.cells, vals):
            dc.set_cell(c, v, size=9.5)
    doc.add_paragraph()


def asset_block(doc, prof, scale, a: dict, photos: Photos, content_w: float) -> None:
    doc.add_heading(f"{a['id']}  {a['name']}", level=3)
    rows = [(prof["scale"]["label"], a["severity"]),
            ("Finding", a["finding"]),
            ("Recorded at", a.get("map_marker") or "no stop"),
            ("Review", "reviewed" if a.get("human_verified") else "unreviewed (auto-accepted)")]
    quantities = a.get("quantities", [])
    if quantities:
        parts = []
        for q in quantities:
            said = f" (said: \"{q['verbatim']}\")" if q.get("verbatim") else ""
            parts.append(open_items.quantity_text(q) + said
                         + ("" if q.get("confirmed") else " " + open_items.QUANTITY_MARKER))
        rows.append(("Quantities", "; ".join(parts)))
    if not a.get("photos"):
        rows.append(("Photos", "none " + open_items.NO_PHOTO_MARKER))
    voice = [e for e in a.get("source_evidence", []) if e.get("type") == "transcript"]
    if voice:
        rows.append(("Voice note", "; ".join(f"{e['ref']}" for e in voice)))
    table = dc.add_table(doc, None, len(rows), 2, [1.4, content_w - 1.4])
    for row, (label, value) in zip(table.rows, rows):
        dc.set_cell(row.cells[0], label, bold=True, size=10, fill=dc.BAND_FILL)
        fill = dc.rating_fill(value, scale) if label == prof["scale"]["label"] else None
        dc.set_cell(row.cells[1], value, size=10, bold=bool(fill), fill=fill)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    photos.grid(a.get("photos", []), content_w)


def add_findings(doc, prof, findings, photos: Photos, headings: list, content_w: float, texts: dict) -> None:
    rep = prof["report"]
    title = "3. " + rep["findings_heading"]
    headings.append((1, title))
    doc.add_heading(title, level=1)
    scale = profiles.scale_values(prof)
    if not prof["classify"]:
        assets = {a["id"]: a for s in findings["sectors"] for a in s.get("assets", [])}
        for st in sorted(findings["stops"], key=lambda s: s["seq"]):
            h = f"{st['title']}"
            headings.append((2, h))
            doc.add_heading(h, level=2)
            dc.add_body(doc, f"{hhmm(st.get('start_ts'))}  |  {len(st['transcripts'])} voice note(s), "
                             f"{len(st['photos'])} photo(s)  |  {'reviewed' if st['human_verified'] else 'unreviewed'}",
                        size=9.5, italic=True)
            for f in st["transcripts"]:
                t = texts.get(f)
                if t:
                    dc.add_body(doc, f"{hhmm(t.get('ts'))}  {t['text_clean']}")
            for aid in st.get("asset_ids", []):
                a = assets.get(aid)
                if not a:
                    continue
                for q in a.get("quantities", []):
                    said = f" (said: \"{q['verbatim']}\")" if q.get("verbatim") else ""
                    dc.add_body(doc, "Quantity: " + open_items.quantity_text(q) + said
                                + ("" if q.get("confirmed") else " " + open_items.QUANTITY_MARKER), size=10)
                if not a.get("photos"):
                    dc.add_body(doc, "Photos: none " + open_items.NO_PHOTO_MARKER, size=10)
            photos.grid(list(dict.fromkeys(st["photos"] + [p for aid in st.get("asset_ids", [])
                                                         for p in (assets.get(aid) or {}).get("photos", [])])),
                        content_w)
        return
    n = 0
    for sector in findings["sectors"]:
        n += 1
        h = f"3.{n} {sector['sector_name']}"
        headings.append((2, h))
        doc.add_heading(h, level=2)
        if not sector.get("present", True):
            dc.add_body(doc, prof["report"].get("absent_finding", "Nothing recorded."), italic=True)
            continue
        dc.add_body(doc, f"{prof['scale']['label']} for this category: {sector['sector_severity']}.",
                    bold=True, size=10.5)
        for a in sector["assets"]:
            asset_block(doc, prof, scale, a, photos, content_w)


def add_priorities(doc, prof, findings, headings: list, content_w: float) -> None:
    title = "4. " + prof["report"]["priorities_heading"]
    headings.append((1, title))
    doc.add_heading(title, level=1)
    rows = findings["priorities"]
    table = dc.add_table(doc, ["Category", "Items", "Action"], len(rows), 3, [1.8, 1.4, content_w - 3.2])
    for row, p in zip(table.rows[1:], rows):
        dc.set_cell(row.cells[0], p["sector_short"], size=10)
        dc.set_cell(row.cells[1], ", ".join(p["asset_ids"]), size=10)
        dc.set_cell(row.cells[2], p["action"], size=10)
    doc.add_paragraph()


def add_open_items(doc, items, headings: list) -> None:
    title = "5. Open items"
    headings.append((1, title))
    doc.add_heading(title, level=1)
    marked = [i for i in items if i["kind"] in ("quantity", "no_photo")]
    gaps = [i for i in items if i["kind"] == "gap"]
    absent = [i for i in items if i["kind"] == "sector_absence"]
    dc.add_body(doc, f"{len(items)} item(s) are listed on the reviewer sheet (review-sheet.docx): "
                     f"{len(marked)} marked in the text, {len(gaps)} missing file(s), "
                     f"{len(absent)} empty categor{'y' if len(absent) == 1 else 'ies'}, "
                     f"{len(items) - len(marked) - len(gaps) - len(absent)} rating(s) to confirm.")
    for g in gaps:
        dc.add_body(doc, "Missing file: " + g["what"])
    for m in marked:
        dc.add_body(doc, f"{m['asset_id']}: {m['what']}", size=10)


def add_voice_log(doc, findings, texts: dict, headings: list) -> None:
    title = "Appendix. Voice-note log"
    headings.append((1, title))
    doc.add_heading(title, level=1)
    order = [f for st in sorted(findings["stops"], key=lambda s: s["seq"]) for f in st["transcripts"]]
    table = dc.add_table(doc, ["File", "Time", "Language", "Text"], len(order), 4, [1.9, 0.6, 0.8, 3.7])
    for row, f in zip(table.rows[1:], order):
        t = texts.get(f, {})
        vals = [f, hhmm(t.get("ts")), t.get("language") or "", t.get("text_clean") or "(not transcribed)"]
        for c, v in zip(row.cells, vals):
            dc.set_cell(c, v, size=9)


# --------------------------------------------------------------------------- build

def build_report(cfg, prof, findings, gaps, run_dir: Path, texts: dict, out_path: Path) -> dict:
    doc = dc.new_document(prof, f"{prof['report']['title']} - {cfg['site']}")
    content_w = dc.content_width_in(doc)
    synthetic = is_synthetic(cfg)
    items = open_items.collect(findings, gaps, prof)
    markers = open_items.report_markers(findings)
    photos = Photos(doc, Path(cfg["media_dir"]), run_dir, photo_captions(findings))
    headings: list[tuple[int, str]] = []

    add_cover(doc, cfg, prof, findings, synthetic)
    toc_anchor = doc.add_paragraph()           # the contents page is spliced in here once headings are known
    add_summary(doc, prof, findings, markers, headings)
    add_stops(doc, cfg, prof, findings, headings, content_w)
    add_findings(doc, prof, findings, photos, headings, content_w, texts)
    add_priorities(doc, prof, findings, headings, content_w)
    add_open_items(doc, items, headings)
    add_voice_log(doc, findings, texts, headings)
    dc.insert_toc(doc, toc_anchor, headings)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    dc.save_docx(doc, out_path)
    return {"photos_embedded": photos.used, "photos_failed": photos.failed, "markers": markers,
            "open_items": len(items), "synthetic_cover": synthetic, "headings": len(headings)}


def _run(cfg, run_dir: Path, log, args) -> int:
    prof = profiles.profile(cfg)
    findings = common.read_stage(run_dir, "findings", "findings")
    if findings.get("profile") != prof["name"]:
        log.error("findings.json was built with profile %r but the config says %r - re-run s6",
                  findings.get("profile"), prof["name"])
        return 5
    gaps = common.read_stage(run_dir, "gaps", "gaps") if (run_dir / "gaps.json").is_file() else None
    texts = transcript_texts(run_dir)
    ups = [run_dir / "findings.json"] + ([run_dir / "gaps.json"] if gaps else [])
    out_path = run_dir / OUT_NAME
    build_path = run_dir / BUILD_NAME
    if (not args.force and out_path.is_file() and build_path.is_file()
            and all(p.stat().st_mtime <= build_path.stat().st_mtime for p in ups)):
        print(f"SKIP {STAGE}: up to date")
        return 0
    info = build_report(cfg, prof, findings, gaps, run_dir, texts, out_path)
    payload = {"status": "partial" if info["photos_failed"] else "ok", "stage": STAGE,
               "profile": prof["name"], "generated_at": datetime.now().strftime("%Y-%m-%dT%H:%M:%S"), **info}
    build_path.write_text(json.dumps(payload, indent=2, sort_keys=True), encoding="utf-8")
    log.info("report: %d photos, %d markers, %d open items -> %s", len(info["photos_embedded"]), info["markers"],
             info["open_items"], out_path)
    return 3 if info["photos_failed"] else 0


def main(argv=None) -> int:
    args = parse_args(argv)
    try:
        cfg = common.load_config(args.config)
    except common.ConfigError as e:
        sys.stderr.write(f"CONFIG ERROR: {e}\n")
        return 2
    run_dir = common.resolve_run_dir(cfg, args.run_dir)
    log = common.get_logger(run_dir, STAGE, args.log_level)
    try:
        return _run(cfg, run_dir, log, args)
    except common.MissingUpstream as e:
        log.error("missing upstream: %s", e)
        return 5
    except Exception:  # noqa: BLE001 - a crash is exit 4 with a traceback in logs/*.err
        tb = traceback.format_exc()
        (run_dir / "logs" / f"{STAGE}.err").write_text(tb, encoding="utf-8")
        log.error("unhandled exception:\n%s", tb)
        return 4


if __name__ == "__main__":
    sys.exit(main())
