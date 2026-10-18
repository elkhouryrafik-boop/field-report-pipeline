#!/usr/bin/env python
"""s10_verify_build.py - stage 10: the deterministic readiness gate over a built packet.

    python -m field_report.stages.s10_verify_build --config <config> [--run-dir DIR]
        [--report report.docx] [--sheet review-sheet.docx] [--quantities quantities.xlsx]

Everything here is mechanically checkable; it does not judge whether a finding is right. Each
check is blocking (exit 4) unless marked advisory:

  G1  the report opens and its body font is the profile's font
  G2  every photo in findings.json is embedded exactly once, and no image is embedded twice
  G3  every number is traced: each quantity's spoken words occur in the transcript evidence cited
      for that asset, and the figure is printed in the report
  G4  every open item is mirrored: the report's [to be confirmed] markers match the unconfirmed
      quantities / photo-less findings, and every item the sheet should carry has its content
      control (tag) in review-sheet.docx; every gaps.json entry reached the sheet
  G5  no named individual in the report, the sheet or the quantities file
  G6  Office metadata is neutral: creator and lastModifiedBy are "field-report-pipeline",
      no company, no template path
  G7  a run with unreviewed stops says "unreviewed" in the report
  G8  quantities.xlsx (profiles that ask for it): one row per quantity, no price/rate/total
      column, unconfirmed rows flagged
  G9  no leftover placeholder text, no personal file path in any rendered text (advisory)

Writes s10_verification.json. Exit codes: 0 pass, 2 bad invocation, 4 blocking problems,
5 a file to check is missing.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import traceback
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from field_report.lib import common, profiles
from field_report.lib import docx_common as dc
from field_report.lib import docx_controls, open_items
from field_report.lib.privacy import name_hits

STAGE = "s10_verify_build"
PRICE_RE = re.compile(r"\b(rate|price|unit price|amount|total|cost)\b", re.I)
PLACEHOLDER_RE = re.compile(r"\[REVIEW\]|\bTODO\b|\bTBD\b|lorem ipsum", re.I)
PERSONAL_PATH_RE = re.compile(r"(?i)\b[a-z]:[\\/]+users[\\/]|/(?:home|users)/[a-z0-9_.-]+/")
ADVISORY = {"PLACEHOLDER_TEXT", "PERSONAL_PATH"}


def problem(code: str, item: str, message: str) -> dict:
    return {"item": item, "code": code, "message": message, "retryable": False, "attempts": 1}


def parse_args(argv=None):
    p = argparse.ArgumentParser(description="stage 10: deterministic readiness gate")
    p.add_argument("--config", required=True)
    p.add_argument("--run-dir", default=None)
    p.add_argument("--report", default=None)
    p.add_argument("--sheet", default=None)
    p.add_argument("--quantities", default=None)
    p.add_argument("--log-level", default="INFO", choices=["DEBUG", "INFO", "WARNING", "ERROR"])
    return p.parse_args(argv)


# --------------------------------------------------------------------------- helpers

def all_photos(findings: dict) -> list[str]:
    seen: list[str] = []
    for s in findings["sectors"]:
        for a in s.get("assets", []):
            seen.extend(a.get("photos", []))
    return seen


def embedded_image_digests(doc) -> list[str]:
    """sha256 of the image behind every inline picture, in document order."""
    out = []
    for blip in doc.element.body.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if rid:
            out.append(hashlib.sha256(doc.part.related_parts[rid].blob).hexdigest())
    return out


def package_metadata(path: Path) -> dict:
    """creator / lastModifiedBy / company / template read straight from the OPC parts."""
    meta = {"creator": None, "lastModifiedBy": None, "company": None, "template": None}
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        if "docProps/core.xml" in names:
            core = z.read("docProps/core.xml").decode("utf-8", "replace")
            for key, tag in (("creator", "dc:creator"), ("lastModifiedBy", "cp:lastModifiedBy")):
                m = re.search(rf"<{tag}[^>]*>([^<]*)</{tag}>", core)
                meta[key] = m.group(1) if m else None
        if "docProps/app.xml" in names:
            app = z.read("docProps/app.xml").decode("utf-8", "replace")
            for key, tag in (("company", "Company"), ("template", "Template")):
                m = re.search(rf"<{tag}[^>]*>([^<]*)</{tag}>", app)
                meta[key] = m.group(1) if m else None
    return meta


def check_metadata(path: Path) -> list[dict]:
    meta = package_metadata(path)
    out = []
    for key in ("creator", "lastModifiedBy"):
        if meta[key] != dc.CREATOR:
            out.append(problem("METADATA_NOT_NEUTRAL", path.name, f"{key} is {meta[key]!r}, expected {dc.CREATOR!r}"))
    for key in ("company", "template"):
        if meta[key]:
            out.append(problem("METADATA_NOT_NEUTRAL", path.name, f"{key} is set ({meta[key]!r})"))
    return out


def text_problems(name: str, text: str, inspectors: list[str]) -> list[dict]:
    out = []
    hits = name_hits(text, inspectors)
    if hits:
        out.append(problem("NAMED_INDIVIDUAL", name, f"{len(hits)} personal name(s) in the text: {hits[:3]}"))
    if PLACEHOLDER_RE.search(text):
        out.append(problem("PLACEHOLDER_TEXT", name, "leftover placeholder text"))
    if PERSONAL_PATH_RE.search(text):
        out.append(problem("PERSONAL_PATH", name, "a personal file path appears in the text"))
    return out


# --------------------------------------------------------------------------- checks

def check_report(path: Path, doc, text: str, findings: dict, cfg: dict, prof: dict, build: dict) -> list[dict]:
    out: list[dict] = []
    # G1
    font = doc.styles["Normal"].font.name
    if font != prof["font"]["name"]:
        out.append(problem("FONT_MISMATCH", path.name, f"body font is {font!r}, profile says {prof['font']['name']!r}"))
    # G2
    photos = all_photos(findings)
    dupes = [p for p, n in Counter(photos).items() if n > 1]
    if dupes:
        out.append(problem("PHOTO_REUSED_IN_FINDINGS", path.name, f"photo(s) assigned to more than one finding: {dupes[:3]}"))
    embedded = build.get("photos_embedded", [])
    missing = sorted(set(photos) - set(embedded))
    if missing:
        out.append(problem("PHOTO_NOT_EMBEDDED", path.name, f"{len(missing)} photo(s) missing from the report: {missing[:3]}"))
    again = [p for p, n in Counter(embedded).items() if n > 1]
    if again:
        out.append(problem("PHOTO_EMBEDDED_TWICE", path.name, f"{again[:3]}"))
    digests = embedded_image_digests(doc)
    expected_images = len(set(embedded)) + (1 if cfg.get("map_image") else 0)
    if len(digests) != expected_images:
        out.append(problem("IMAGE_COUNT", path.name, f"{len(digests)} images in the document, expected {expected_images}"))
    if len(set(digests)) != len(digests):
        out.append(problem("DUPLICATE_IMAGE", path.name, "the same image is embedded more than once"))
    # G3
    flat = re.sub(r"\s+", " ", text)
    for s in findings["sectors"]:
        for a in s.get("assets", []):
            spoken = " ".join(e.get("quote", "") for e in a.get("source_evidence", []) if e.get("type") == "transcript").lower()
            for i, q in enumerate(a.get("quantities", [])):
                where = f"{a['id']}/qty/{i}"
                if q.get("confirmed") and any(e.get("type") == "reviewer" for e in a.get("source_evidence", [])):
                    pass  # confirmed by a person: the reviewer is the evidence
                elif not q.get("verbatim") or q["verbatim"].lower() not in spoken:
                    out.append(problem("UNTRACED_NUMBER", where, "the words behind this figure are not in the cited transcript"))
                if open_items.quantity_text(q) not in flat:
                    out.append(problem("NUMBER_NOT_IN_REPORT", where, f"{open_items.quantity_text(q)!r} is not printed in the report"))
    # G4 (report side)
    n_markers = len(re.findall(r"\[to be confirmed[^\]]*\]", text))
    expected = open_items.report_markers(findings)
    if n_markers != expected:
        out.append(problem("MARKER_COUNT", path.name, f"{n_markers} [to be confirmed] marker(s) in the text, expected {expected}"))
    # G7
    if any(not st.get("human_verified") for st in findings.get("stops", [])) and "nreviewed" not in text:
        out.append(problem("UNREVIEWED_NOT_STATED", path.name, "stops are unreviewed but the report does not say so"))
    if build.get("synthetic_cover") and "synthetic sample data" not in text:
        out.append(problem("SYNTHETIC_NOT_STATED", path.name, "sample data without the synthetic cover line"))
    return out


def check_sheet(path: Path, sheet_json: dict, findings: dict, gaps: dict | None, prof: dict) -> list[dict]:
    out: list[dict] = []
    expected = open_items.collect(findings, gaps, prof)
    want = {i["key"] for i in expected}
    have = {i["key"] for i in sheet_json.get("items", [])}
    if want - have:
        out.append(problem("OPEN_ITEM_NOT_ON_SHEET", path.name, f"{len(want - have)} open item(s) missing from confirmation_sheet.json: {sorted(want - have)[:3]}"))
    if have - want:
        out.append(problem("SHEET_ITEM_STALE", path.name, f"{len(have - want)} sheet item(s) no longer open: {sorted(have - want)[:3]}"))
    tags = set(docx_controls.read_controls(path))
    for item in sheet_json.get("items", []):
        if item["key"] not in tags:
            out.append(problem("SHEET_CONTROL_MISSING", item["key"], "no dropdown control with this tag in review-sheet.docx"))
        elif item["kind"] != "severity" and f"{item['key']}#correction" not in tags:
            out.append(problem("SHEET_CONTROL_MISSING", item["key"], "no correction box with this tag"))
    for g in (gaps or {}).get("gaps", []):
        if f"gap/{g.get('message_id') or g.get('filename')}" not in have:
            out.append(problem("GAP_NOT_ON_SHEET", str(g.get("message_id")), "a media gap did not reach the sheet"))
    return out


def check_quantities(path: Path, findings: dict) -> list[dict]:
    from openpyxl import load_workbook
    out: list[dict] = []
    wb = load_workbook(path)
    ws = wb["Quantities"]
    header = [str(c.value or "") for c in ws[1]]
    bad = [h for h in header if PRICE_RE.search(h)]
    if bad:
        out.append(problem("PRICE_COLUMN", path.name, f"price-like column header(s): {bad}"))
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    n_expected = sum(len(a.get("quantities", [])) for s in findings["sectors"] if s.get("present", True) for a in s["assets"])
    if len(rows) != n_expected:
        out.append(problem("QUANTITY_ROWS", path.name, f"{len(rows)} row(s), findings.json has {n_expected} quantities"))
    n_open = sum(1 for s in findings["sectors"] if s.get("present", True) for a in s["assets"]
                 for q in a.get("quantities", []) if not q.get("confirmed"))
    flagged = sum(1 for r in rows if r[6] == "To be confirmed")
    if flagged != n_open:
        out.append(problem("QUANTITY_STATUS", path.name, f"{flagged} row(s) flagged to be confirmed, expected {n_open}"))
    return out


def review_coverage(findings: dict) -> dict:
    stops = findings.get("stops", [])
    done = sum(1 for s in stops if s.get("human_verified"))
    return {"verified": done, "total": len(stops), "pct": round(100 * done / len(stops)) if stops else 0}


# --------------------------------------------------------------------------- main

def _run(cfg, run_dir: Path, log, args) -> int:
    prof = profiles.profile(cfg)
    findings = common.read_stage(run_dir, "findings", "findings")
    gaps = common.read_stage(run_dir, "gaps", "gaps") if (run_dir / "gaps.json").is_file() else None
    sheet_json = common.read_stage(run_dir, "confirmation_sheet", "confirmation_sheet")
    report = Path(args.report) if args.report else run_dir / "report.docx"
    sheet = Path(args.sheet) if args.sheet else run_dir / "review-sheet.docx"
    qpath = Path(args.quantities) if args.quantities else run_dir / "quantities.xlsx"
    need = [report, sheet] + ([qpath] if prof.get("quantities_sheet") else [])
    for p in need:
        if not p.is_file():
            raise common.MissingUpstream(f"file to check is missing: {p}")
    build_path = run_dir / "report_build.json"
    build = json.loads(build_path.read_text(encoding="utf-8")) if build_path.is_file() else {}
    if not build:
        raise common.MissingUpstream(f"missing {build_path} - run s7 first")

    inspectors = cfg.get("inspectors", [])
    problems: list[dict] = []
    doc = Document(str(report))
    text = dc.docx_text(doc)
    problems += check_report(report, doc, text, findings, cfg, prof, build)
    problems += check_metadata(report)
    problems += text_problems(report.name, text, inspectors)
    problems += check_sheet(sheet, sheet_json, findings, gaps, prof)
    problems += check_metadata(sheet)
    problems += text_problems(sheet.name, dc.docx_text(Document(str(sheet))), inspectors)
    if prof.get("quantities_sheet"):
        problems += check_quantities(qpath, findings)
        problems += check_metadata(qpath)
        from openpyxl import load_workbook
        cells = " ".join(str(c.value) for ws in load_workbook(qpath) for row in ws.iter_rows() for c in row if c.value)
        problems += text_problems(qpath.name, cells, inspectors)

    blocking = [p for p in problems if p["code"] not in ADVISORY]
    payload = common.envelope(STAGE, cfg["site"], cfg,
                              {"assets": sum(len(s.get("assets", [])) for s in findings["sectors"])},
                              {"problems": len(problems), "blocking": len(blocking)})
    payload["errors"] = problems
    payload["checked"] = {"report": report.name, "sheet": sheet.name,
                          "quantities": qpath.name if prof.get("quantities_sheet") else None}
    payload["status"] = "failed" if blocking else ("partial" if problems else "ok")
    payload["review_coverage"] = review_coverage(findings)
    out = run_dir / "s10_verification.json"
    tmp = out.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")
    os.replace(tmp, out)

    for p in problems:
        (log.warning if p["code"] in ADVISORY else log.error)("%s [%s] %s", p["code"], p["item"], p["message"])
    rc = payload["review_coverage"]
    print(f"review coverage: {rc['verified']}/{rc['total']} stops reviewed ({rc['pct']}%)")
    if blocking:
        print(f"GATE FAILED: {len(blocking)} blocking problem(s)", file=sys.stderr)
        return 4
    print(f"GATE PASSED: {len(problems)} advisory note(s)")
    return 0


def main(argv=None) -> int:
    args = parse_args(argv)
    try:
        cfg = common.load_config(args.config)
    except common.ConfigError as e:
        sys.stderr.write(f"CONFIG ERROR: {e}\n")
        return 2
    run_dir = common.resolve_run_dir(cfg, args.run_dir)
    log = common.get_logger(run_dir, STAGE, args.log_level)
    try:
        return _run(cfg, run_dir, log, args)
    except common.MissingUpstream as e:
        log.error("missing upstream: %s", e)
        return 5
    except Exception:  # noqa: BLE001
        tb = traceback.format_exc()
        (run_dir / "logs" / f"{STAGE}.err").write_text(tb, encoding="utf-8")
        log.error("unhandled exception:\n%s", tb)
        return 4


if __name__ == "__main__":
    sys.exit(main())
