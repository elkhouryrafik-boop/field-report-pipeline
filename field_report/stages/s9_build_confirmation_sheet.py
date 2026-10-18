#!/usr/bin/env python
"""s9_build_confirmation_sheet.py - stage 9: open items -> review-sheet.docx (Word content controls).

    python -m field_report.stages.s9_build_confirmation_sheet --config <config> [--run-dir DIR] [--revision Rev1]

The reviewer sheet lists every open item (lib/open_items.py - the same list the report brackets
and the gate checks). The reviewer clicks a dropdown per row ("Yes - as recorded" / "No - see
correction" / "Don't know yet", or a rating), and types only when correcting a figure. Each control
carries the item key as its tag, so the answers survive sorting or deleting rows:

    python -m field_report ingest-review --config <config> --from-docx review-sheet.docx

Writes review-sheet.docx and confirmation_sheet.json (the contract ingestion reads).
Exit codes: 0 ok, 2 bad invocation / file locked, 4 failed, 5 missing upstream.
"""
from __future__ import annotations

import argparse
import sys
import traceback
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH

from field_report.lib import common, profiles
from field_report.lib import docx_common as dc
from field_report.lib import docx_controls as controls
from field_report.lib import open_items

STAGE = "s9_build_confirmation_sheet"
OUT_NAME = "review-sheet.docx"
HEADER = ["#", "Where", "Question", "Why it is open", "Answer", "Correction (only if No)", "Source"]
COL_WIDTHS = [0.35, 1.45, 2.2, 2.1, 1.3, 1.35, 1.05]
RATING_HEADER = ["#", "Where", "Current", "What the rating rests on", "Change to"]
RATING_WIDTHS = [0.35, 2.1, 1.0, 4.7, 1.7]
ANSWER_FILL, CORRECTION_FILL = "FFF2A8", "FFFDE7"


def parse_args(argv=None):
    p = argparse.ArgumentParser(description="stage 9: build the reviewer sheet")
    p.add_argument("--config", required=True)
    p.add_argument("--run-dir", default=None)
    p.add_argument("--revision", default="Rev1")
    p.add_argument("--out", default=None)
    p.add_argument("--force", action="store_true")
    p.add_argument("--log-level", default="INFO", choices=["DEBUG", "INFO", "WARNING", "ERROR"])
    return p.parse_args(argv)


def where(item: dict) -> str:
    if item.get("asset_id"):
        return f"{item['asset_id']} {item['asset_name']}"
    return item.get("sector_name") or "Whole report"


def build_docx(items: list[dict], cfg: dict, prof: dict, revision: str, synthetic: bool):
    doc = dc.new_document(prof, f"Review sheet {revision} - {cfg['site']}", landscape=True, margin_in=0.6)
    dc.add_body(doc, f"{prof['report']['title']}: review sheet {revision}", bold=True, size=15,
                align=WD_ALIGN_PARAGRAPH.CENTER, color=dc.HEADING_COLOR)
    dc.add_body(doc, profiles.locality_line(cfg), bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    if synthetic:
        dc.add_body(doc, "Generated from synthetic sample data", bold=True, size=10,
                    align=WD_ALIGN_PARAGRAPH.CENTER, color=dc.MARKER_COLOR)
    doc.add_heading("How to fill this in", level=1)
    dc.add_body(doc, "Every row is a point in the report that the evidence could not settle, with the reason it "
                     "is still open. Click the yellow cell and pick an answer; there is nothing to type unless "
                     "you are correcting a figure.")
    dc.add_body(doc, "Yes - as recorded: the report keeps the figure and the item stops being flagged.")
    dc.add_body(doc, "No - see correction: type the right figure in the next column, nothing else.")
    dc.add_body(doc, "Don't know yet: the item stays open and comes back on the next revision.")

    plain = [i for i in items if i["kind"] != "severity"]
    ratings = [i for i in items if i["kind"] == "severity"]
    doc.add_heading("Open items", level=1)
    if plain:
        table = dc.add_table(doc, HEADER, len(plain), len(HEADER), COL_WIDTHS, size=9.5)
        for row, item in zip(table.rows[1:], plain):
            c = row.cells
            fill = dc.BAND_FILL if item["item_no"] % 2 else None
            dc.set_cell(c[0], str(item["item_no"]), bold=True, size=9.5, fill=fill)
            dc.set_cell(c[1], where(item), color=dc.HEADING_COLOR, size=9.5, fill=fill)
            dc.set_cell(c[2], item["question"], size=9.5, fill=fill)
            dc.set_cell(c[3], item["why"], size=9, fill=fill)
            dc.set_cell(c[6], item["source"], size=9, fill=fill)
            dc.shade_cell(c[4], ANSWER_FILL)
            dc.shade_cell(c[5], CORRECTION_FILL)
            controls.insert_dropdown(c[4], item["key"], open_items.options_for(item, prof), alias="Answer",
                                     selected=open_items.CHOOSE)
            controls.insert_text_box(c[5], f"{item['key']}#correction")
    else:
        dc.add_body(doc, "None: nothing in the report is open.", italic=True)

    if ratings:
        label = prof["scale"]["label"]
        doc.add_heading(f"{label} ratings", level=1)
        dc.add_body(doc, f"Scale: {', '.join(prof['scale']['values'])}. Each category takes the rating of its "
                         f"worst item, so a change here re-rates the category and the summary. Each row shows "
                         f"what the rating currently rests on.")
        table = dc.add_table(doc, RATING_HEADER, len(ratings), len(RATING_HEADER), RATING_WIDTHS, size=9.5)
        for row, item in zip(table.rows[1:], ratings):
            c = row.cells
            fill = dc.BAND_FILL if item["item_no"] % 2 else None
            dc.set_cell(c[0], str(item["item_no"]), bold=True, size=9.5, fill=fill)
            dc.set_cell(c[1], where(item), color=dc.HEADING_COLOR, size=9.5, fill=fill)
            dc.set_cell(c[2], item["current_value"] or "", bold=True, size=9.5, fill=fill)
            dc.set_cell(c[3], item["why"], size=9, fill=fill)
            dc.shade_cell(c[4], ANSWER_FILL)
            controls.insert_dropdown(c[4], item["key"], open_items.options_for(item, prof), alias=label,
                                     selected=open_items.KEEP_OPTION)
    dc.add_body(doc, "")
    counts: dict[str, int] = {}
    for it in items:
        counts[it["kind"]] = counts.get(it["kind"], 0) + 1
    listed = ", ".join(f"{n} {k.replace('_', ' ')}" for k, n in sorted(counts.items()))
    dc.add_body(doc, f"{len(items)} item(s){': ' + listed if listed else ''}.", italic=True)
    return doc


def _run(cfg, run_dir: Path, log, args) -> int:
    prof = profiles.profile(cfg)
    findings = common.read_stage(run_dir, "findings", "findings")
    gaps = common.read_stage(run_dir, "gaps", "gaps") if (run_dir / "gaps.json").is_file() else None
    items = open_items.collect(findings, gaps, prof)
    from field_report.stages.s7_build_report import is_synthetic
    out_path = Path(args.out) if args.out else run_dir / OUT_NAME
    n_assets = sum(len(s["assets"]) for s in findings["sectors"] if s.get("present", True))
    payload = common.envelope(STAGE, cfg["site"], cfg, {"assets": n_assets}, {"items": len(items)})
    payload.update({"site": cfg["site"], "revision": args.revision, "items": items, "out_path": str(out_path)})
    doc = build_docx(items, cfg, prof, args.revision, is_synthetic(cfg))
    try:
        out_path.parent.mkdir(parents=True, exist_ok=True)
        dc.save_docx(doc, out_path)
    except PermissionError:
        print(f"LOCKED: {out_path} is open in another program; close it or pass --out", file=sys.stderr)
        return 2
    common.write_stage(run_dir, "confirmation_sheet", payload, "confirmation_sheet")
    log.info("review sheet: %d item(s) -> %s", len(items), out_path)
    return 0


def main(argv=None) -> int:
    args = parse_args(argv)
    try:
        cfg = common.load_config(args.config)
    except common.ConfigError as e:
        sys.stderr.write(f"CONFIG ERROR: {e}\n")
        return 2
    run_dir = common.resolve_run_dir(cfg, args.run_dir)
    log = common.get_logger(run_dir, STAGE, args.log_level)
    try:
        return _run(cfg, run_dir, log, args)
    except common.MissingUpstream as e:
        log.error("missing upstream: %s", e)
        return 5
    except Exception:  # noqa: BLE001
        tb = traceback.format_exc()
        (run_dir / "logs" / f"{STAGE}.err").write_text(tb, encoding="utf-8")
        log.error("unhandled exception:\n%s", tb)
        return 4


if __name__ == "__main__":
    sys.exit(main())
