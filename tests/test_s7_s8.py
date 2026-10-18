"""s7 (report.docx) and s8 (quantities.xlsx) on the synthetic sample."""
from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from field_report import demo
from field_report.lib import docx_common as dc
from field_report.stages import s7_build_report as s7
from field_report.stages import s8_build_quantities as s8


def load(rd, name):
    return json.loads((rd / f"{name}.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def doc(built_run):
    return Document(str(built_run[1] / "report.docx"))


@pytest.fixture(scope="module")
def text(doc):
    return dc.docx_text(doc)


def test_report_exists_and_opens(built_run):
    assert (built_run[1] / "report.docx").stat().st_size > 5000


def test_report_mentions_the_site(built_run, text):
    cfg = json.loads(built_run[0].read_text(encoding="utf-8"))
    assert cfg["site"] in text


def test_report_has_a_synthetic_banner(text):
    assert "synthetic" in text.lower()


def test_every_photo_embedded_exactly_once(built_run):
    rd = built_run[1]
    build = load(rd, "report_build")
    findings = load(rd, "findings")
    in_findings = [p for s in findings["sectors"] for a in s["assets"] for p in a["photos"]]
    assert len(build["photos_embedded"]) == len(set(build["photos_embedded"])) == len(in_findings)
    assert build["photos_failed"] == []


def test_embedded_images_are_distinct(built_run):
    with zipfile.ZipFile(built_run[1] / "report.docx") as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
    assert len(media) == 13          # 12 photos + the stop map


def test_unreviewed_run_says_unreviewed(text):
    assert "unreviewed" in text.lower()


def test_to_be_confirmed_markers_match_build_record(built_run, text):
    n = len(re.findall(r"\[to be (?:confirmed|verified)", text))
    assert n == load(built_run[1], "report_build")["markers"] > 0


def test_markers_are_bold_runs(doc):
    runs = [r for p in doc.paragraphs for r in p.runs if dc.MARKER_RE.fullmatch(r.text)]
    runs += [r for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs
             if dc.MARKER_RE.fullmatch(r.text)]
    assert runs
    assert all(r.bold for r in runs)


def test_core_properties_are_neutral(doc):
    cp = doc.core_properties
    assert cp.author == cp.last_modified_by == "field-report-pipeline"


def test_app_xml_has_no_template_or_company(built_run):
    with zipfile.ZipFile(built_run[1] / "report.docx") as z:
        app = z.read("docProps/app.xml").decode()
    assert "Company" not in app and ".dotm" not in app


def test_build_record_is_ok(built_run):
    b = load(built_run[1], "report_build")
    assert b["status"] == "ok" and b["profile"] == "site-inspection"


def test_s7_rerun_is_idempotent_skip(fresh_run, capsys):
    cfg, rd = fresh_run
    before = (rd / "report.docx").stat().st_mtime_ns
    assert s7.main(["--config", str(cfg), "--run-dir", str(rd)]) == 0
    assert "up to date" in capsys.readouterr().out
    assert (rd / "report.docx").stat().st_mtime_ns == before


def test_s7_force_rebuilds(fresh_run):
    cfg, rd = fresh_run
    assert s7.main(["--config", str(cfg), "--run-dir", str(rd), "--force"]) == 0
    assert (rd / "report.docx").is_file()


def test_s7_wrong_profile_exits_5(fresh_run):
    cfg, rd = fresh_run
    f = load(rd, "findings")
    f["profile"] = "event-log"
    (rd / "findings.json").write_text(json.dumps(f), encoding="utf-8")
    assert s7.main(["--config", str(cfg), "--run-dir", str(rd), "--force"]) == 5


def test_s7_missing_findings_exits_5(fresh_run):
    cfg, rd = fresh_run
    (rd / "findings.json").unlink()
    assert s7.main(["--config", str(cfg), "--run-dir", str(rd), "--force"]) == 5


def test_s7_bad_config_exits_2(tmp_path):
    assert s7.main(["--config", str(tmp_path / "nope.json")]) == 2


def test_a_missing_photo_gives_exit_3_and_a_placeholder(fresh_run):
    cfg, rd = fresh_run
    c = json.loads(cfg.read_text(encoding="utf-8"))
    media = Path(c["media_dir"])
    if not media.is_absolute():
        media = cfg.parent / media
    findings = load(rd, "findings")
    first = next(p for s in findings["sectors"] for a in s["assets"] for p in a["photos"])
    (media / first).unlink()
    assert s7.main(["--config", str(cfg), "--run-dir", str(rd), "--force"]) == 3
    assert load(rd, "report_build")["status"] == "partial"
    assert "could not be embedded" in dc.docx_text(Document(str(rd / "report.docx")))


@pytest.mark.parametrize("iso,expected", [("2030-01-14", "14 January 2030"),
                                          ("2030-12-01T10:00:00", "1 December 2030")])
def test_long_date(iso, expected):
    assert s7.long_date(iso) == expected


@pytest.mark.parametrize("ts,expected", [("2030-01-14T09:05:00", "09:05"), (None, ""), ("", "")])
def test_hhmm(ts, expected):
    assert s7.hhmm(ts) == expected


def _stops(done, total):
    return {"stops": [{"human_verified": i < done} for i in range(total)]}


def test_review_sentence_all():
    assert s7.review_sentence(_stops(3, 3)).startswith("All 3")


def test_review_sentence_none():
    assert s7.review_sentence(_stops(0, 3)).startswith("Unreviewed")


def test_review_sentence_some():
    assert "1 of 3" in s7.review_sentence(_stops(1, 3))


def test_is_synthetic(tmp_path):
    assert not s7.is_synthetic({"_dir": str(tmp_path)})
    (tmp_path / "sample_manifest.json").write_text("{}")
    assert s7.is_synthetic({"_dir": str(tmp_path)})


# --- s8 --------------------------------------------------------------------------------

@pytest.fixture(scope="module")
def damage(tmp_path_factory):
    out = tmp_path_factory.mktemp("damage") / "o"
    assert demo.main(["--offline", "--profile", "damage-assessment", "--out", str(out)]) == 0
    return out


def test_damage_demo_has_quantities_workbook(damage):
    assert (damage / "quantities.xlsx").is_file()


def test_site_inspection_has_no_quantities_workbook(built_run):
    assert not (built_run[1] / "quantities.xlsx").exists()


def test_workbook_rows_match_findings(damage):
    findings = load(damage / "work", "findings")
    expected = sum(len(a.get("quantities", [])) for s in findings["sectors"] if s.get("present", True)
                   for a in s["assets"])
    assert len(s8.build_rows(findings)) == expected > 0


def test_workbook_has_no_price_columns(damage):
    wb = load_workbook(damage / "quantities.xlsx")
    heads = " ".join(str(c.value) for ws in wb for row in ws.iter_rows(max_row=8) for c in row if c.value)
    assert not re.search(r"\b(rate|price|total|cost|amount)\b", heads, re.I)


def test_workbook_metadata_is_neutral(damage):
    wb = load_workbook(damage / "quantities.xlsx")
    assert wb.properties.creator == wb.properties.lastModifiedBy == "field-report-pipeline"


def test_build_rows_confirmed_is_boolean(damage):
    rows = s8.build_rows(load(damage / "work", "findings"))
    assert all(isinstance(r["confirmed"], bool) for r in rows)
    assert all(r["unit"] != "other" for r in rows)


def test_build_rows_skips_absent_sectors():
    f = {"sectors": [{"present": False, "sector_name": "X", "assets": [
        {"id": "A1", "name": "n", "quantities": [{"unit": "m", "value": 1, "confirmed": False}]}]}]}
    assert s8.build_rows(f) == []


def test_build_rows_unit_other_is_blank():
    f = {"sectors": [{"present": True, "sector_name": "X", "assets": [
        {"id": "A1", "name": "n", "map_marker": "3", "source_evidence": [{"type": "transcript", "ref": "PTT-1"}],
         "quantities": [{"unit": "other", "value": 2, "confirmed": True, "verbatim": "two"}]}]}]}
    (row,) = s8.build_rows(f)
    assert row["unit"] == "" and row["number"] == 2 and row["source"] == "PTT-1" and row["stop"] == "3"
