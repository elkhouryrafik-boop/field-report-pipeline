"""lib/docx_common.py: neutral metadata, marker runs, images, helpers."""
from __future__ import annotations

import zipfile

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image

from field_report.lib import docx_common as dc
from field_report.lib import profiles


@pytest.fixture(scope="module")
def prof():
    return profiles.load("site-inspection")


def test_new_document_has_neutral_properties(prof):
    cp = dc.new_document(prof, "T").core_properties
    assert cp.author == cp.last_modified_by == "field-report-pipeline"
    assert cp.title == "T"
    assert cp.comments == "" and cp.keywords == ""


def test_new_document_uses_the_profile_font(prof):
    doc = dc.new_document(prof, "T")
    assert doc.styles["Normal"].font.name == prof["font"]["name"]
    assert doc.styles["Heading 1"].font.name == prof["font"]["name"]


@pytest.mark.parametrize("landscape,w,h", [(False, 8.5, 11), (True, 11, 8.5)])
def test_orientation(prof, landscape, w, h):
    sec = dc.new_document(prof, "T", landscape=landscape).sections[0]
    assert (sec.page_width.inches, sec.page_height.inches) == (w, h)


def test_content_width(prof):
    assert dc.content_width_in(dc.new_document(prof, "T", margin_in=1.0)) == pytest.approx(6.5)


def test_save_docx_scrubs_app_xml(prof, tmp_path):
    path = tmp_path / "x.docx"
    dc.save_docx(dc.new_document(prof, "T"), path)
    with zipfile.ZipFile(path) as z:
        app = z.read("docProps/app.xml").decode()
    assert "field-report-pipeline" in app
    assert "Normal.dotm" not in app and "Company" not in app
    assert not path.with_suffix(".tmp").exists()


def test_saved_docx_reopens(prof, tmp_path):
    path = tmp_path / "x.docx"
    doc = dc.new_document(prof, "T")
    dc.add_body(doc, "hello")
    dc.save_docx(doc, path)
    assert "hello" in dc.docx_text(Document(str(path)))


def test_workbook_properties_are_neutral():
    wb = Workbook()
    dc.neutral_workbook_properties(wb, "Q")
    assert wb.properties.creator == wb.properties.lastModifiedBy == "field-report-pipeline"
    assert wb.properties.title == "Q"


@pytest.mark.parametrize("text,n_markers", [
    ("plain", 0),
    ("12 m [to be confirmed]", 1),
    ("[to be confirmed] and [to be verified: depth]", 2),
])
def test_marked_runs(prof, text, n_markers):
    doc = dc.new_document(prof, "T")
    p = doc.add_paragraph()
    dc.write_marked_runs(p, text)
    marked = [r for r in p.runs if dc.MARKER_RE.fullmatch(r.text)]
    assert len(marked) == n_markers
    assert all(r.bold and r.font.color.rgb is not None for r in marked)
    assert "".join(r.text for r in p.runs) == text


def test_set_cell_resets_previous_text(prof):
    doc = dc.new_document(prof, "T")
    t = doc.add_table(rows=1, cols=1)
    dc.set_cell(t.cell(0, 0), "one")
    dc.set_cell(t.cell(0, 0), "two", bold=True)
    assert t.cell(0, 0).text == "two"


def test_add_table_has_a_repeating_header(prof):
    doc = dc.new_document(prof, "T")
    t = dc.add_table(doc, ["A", "B"], 3, 2, [1.0, 2.0])
    assert len(t.rows) == 4
    assert "w:tblHeader" in t.rows[0]._tr.xml
    assert t.rows[0].cells[0].text == "A"


def test_add_table_without_header(prof):
    t = dc.add_table(dc.new_document(prof, "T"), None, 2, 3)
    assert len(t.rows) == 2


@pytest.mark.parametrize("value,expected_index", [("High", 0), ("Medium", 1), ("Low", 2)])
def test_rating_fill_follows_rank(value, expected_index):
    scale = ["High", "Medium", "Low"]
    assert dc.rating_fill(value, scale) == dc.RANK_FILLS[expected_index]


def test_rating_fill_unknown_is_grey():
    assert dc.rating_fill("N/A", ["High", "Low"]) == dc.NA_FILL


def test_page_break_adds_a_break(prof):
    doc = dc.new_document(prof, "T")
    dc.page_break(doc)
    assert 'w:br w:type="page"' in doc.element.body.xml or "w:type=\"page\"" in doc.element.body.xml


def test_toc_lists_every_heading(prof):
    doc = dc.new_document(prof, "T")
    anchor = doc.add_paragraph()
    dc.insert_toc(doc, anchor, [(1, "Summary"), (2, "Detail")])
    text = dc.docx_text(doc)
    assert "Summary" in text and "Detail" in text


def test_photo_stream_strips_exif_and_bounds_size(tmp_path):
    src = tmp_path / "p.jpg"
    im = Image.new("RGB", (3000, 2000), "red")
    exif = Image.Exif()
    exif[0x010F] = "SecretCam"
    im.save(src, exif=exif)
    out = Image.open(dc.photo_stream(src, max_px=800))
    assert max(out.size) <= 800
    assert not out.getexif()


def test_photo_stream_converts_png(tmp_path):
    src = tmp_path / "p.png"
    Image.new("RGBA", (50, 40), (0, 0, 0, 0)).save(src)
    assert Image.open(dc.photo_stream(src)).mode == "RGB"
